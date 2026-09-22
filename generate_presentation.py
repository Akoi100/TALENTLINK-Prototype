from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE

def create_rich_presentation():
    prs = Presentation()
    # 16:9 aspect ratio
    prs.slide_width = Inches(13.33)
    prs.slide_height = Inches(7.5)

    # Colors
    bg_dark = RGBColor(13, 71, 161) # Deep Royal Blue
    bg_light = RGBColor(248, 250, 252) # Slate-50
    accent_gold = RGBColor(255, 179, 0)
    text_dark = RGBColor(30, 41, 59)
    white = RGBColor(255, 255, 255)

    def set_bg(slide, color):
        background = slide.background
        fill = background.fill
        fill.solid()
        fill.fore_color.rgb = color

    def add_title(slide, text, color):
        txBox = slide.shapes.add_textbox(Inches(1), Inches(0.5), Inches(11.33), Inches(1))
        p = txBox.text_frame.paragraphs[0]
        p.text = text
        p.font.size = Pt(40)
        p.font.bold = True
        p.font.color.rgb = color
        return txBox

    # [SLIDES 1-8 CODE PRESERVED]
    
    # ==========================================
    # Slide 1: Title Slide
    # ==========================================
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_bg(slide, bg_dark)
    
    txBox = slide.shapes.add_textbox(Inches(1), Inches(2.5), Inches(11.33), Inches(1.5))
    p = txBox.text_frame.paragraphs[0]
    p.text = "TALENTLINK"
    p.font.size = Pt(72)
    p.font.bold = True
    p.font.color.rgb = white
    p.alignment = PP_ALIGN.CENTER
    
    txBox2 = slide.shapes.add_textbox(Inches(1), Inches(4), Inches(11.33), Inches(1))
    p2 = txBox2.text_frame.paragraphs[0]
    p2.text = "Official Co-Curricular Tracking & Talent Scouting Platform"
    p2.font.size = Pt(28)
    p2.font.color.rgb = accent_gold
    p2.alignment = PP_ALIGN.CENTER

    # ==========================================
    # Slide 2: The Problem
    # ==========================================
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_bg(slide, bg_light)
    add_title(slide, "The Problem: Fragmented Ecosystems", bg_dark)

    problems = [
        ("No Official Verification", "Achievements are self-reported. Hard for sponsors to trust data."),
        ("Disconnected Opportunities", "Scouts and sponsors cannot easily find rising local talent."),
        ("No Centralized Record", "Co-curricular data isn't tracked officially like academics.")
    ]
    
    for i, (title, desc) in enumerate(problems):
        shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1 + (i * 3.8)), Inches(2.5), Inches(3.5), Inches(3.5))
        shape.fill.solid()
        shape.fill.fore_color.rgb = white
        shape.line.color.rgb = accent_gold
        shape.line.width = Pt(2)
        
        tf = shape.text_frame
        tf.word_wrap = True
        p1 = tf.paragraphs[0]
        p1.text = f"\n{title}\n"
        p1.font.bold = True
        p1.font.size = Pt(24)
        p1.font.color.rgb = bg_dark
        p1.alignment = PP_ALIGN.CENTER
        
        p2 = tf.add_paragraph()
        p2.text = desc
        p2.font.size = Pt(18)
        p2.font.color.rgb = text_dark
        p2.alignment = PP_ALIGN.CENTER

    # ==========================================
    # Slide 3: The Solution
    # ==========================================
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_bg(slide, bg_light)
    add_title(slide, "The Solution: A Single Verifiable Hub", bg_dark)
    
    blocks = [
        ("1. Certified Input", "Officials input scores directly from the field.", Inches(1)),
        ("2. Student Dossier", "System aggregates verified metrics & badges.", Inches(5)),
        ("3. Scouting & Offers", "Sponsors filter data & offer scholarships.", Inches(9))
    ]
    
    for i, (title, desc, left) in enumerate(blocks):
        shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, Inches(3), Inches(3.2), Inches(2.5))
        shape.fill.solid()
        shape.fill.fore_color.rgb = bg_dark
        
        tf = shape.text_frame
        tf.word_wrap = True
        p1 = tf.paragraphs[0]
        p1.text = f"\n{title}\n"
        p1.font.bold = True
        p1.font.size = Pt(22)
        p1.font.color.rgb = accent_gold
        p1.alignment = PP_ALIGN.CENTER
        
        p2 = tf.add_paragraph()
        p2.text = desc
        p2.font.size = Pt(16)
        p2.font.color.rgb = white
        p2.alignment = PP_ALIGN.CENTER
        
        if i < 2:
            arrow = slide.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, left + Inches(3.3), Inches(4), Inches(0.6), Inches(0.4))
            arrow.fill.solid()
            arrow.fill.fore_color.rgb = accent_gold

    # ==========================================
    # Slide 4: Target Personas
    # ==========================================
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_bg(slide, bg_dark)
    add_title(slide, "Who is this for?", white)
    
    personas = [
        ("Referees & Judges", "Provide the official data."),
        ("Scouts & Sponsors", "Consume data to offer opportunities."),
        ("Students", "Build a verifiable track record.")
    ]
    
    for i, (role, desc) in enumerate(personas):
        shape = slide.shapes.add_shape(MSO_SHAPE.SNIP_2_DIAG_RECTANGLE, Inches(1 + (i*4)), Inches(2.5), Inches(3.5), Inches(3.5))
        shape.fill.solid()
        shape.fill.fore_color.rgb = white
        
        tf = shape.text_frame
        tf.word_wrap = True
        p1 = tf.paragraphs[0]
        p1.text = f"\n\n{role}\n"
        p1.font.bold = True
        p1.font.size = Pt(24)
        p1.font.color.rgb = bg_dark
        p1.alignment = PP_ALIGN.CENTER
        
        p2 = tf.add_paragraph()
        p2.text = desc
        p2.font.size = Pt(18)
        p2.font.color.rgb = text_dark
        p2.alignment = PP_ALIGN.CENTER

    # ==========================================
    # Slide 5: Core Process 1 - Verification
    # ==========================================
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_bg(slide, bg_light)
    add_title(slide, "Core Process 1: Verification & Scoring", bg_dark)

    steps = [
        ("1. Event Occurs", "A sports match, play, or debate happens.", Inches(1), Inches(3)),
        ("2. Official Login", "Referee/Adjudicator logs into workspace.", Inches(5), Inches(2)),
        ("3. Submits Data", "Scores & MVP candidates are entered.", Inches(9), Inches(3)),
        ("4. System Lock", "Data becomes immutable & updates Student Profile.", Inches(5), Inches(5))
    ]
    
    for title, desc, left, top in steps:
        shape = slide.shapes.add_shape(MSO_SHAPE.PENTAGON, left, top, Inches(3.2), Inches(1.8))
        shape.fill.solid()
        shape.fill.fore_color.rgb = bg_dark
        shape.line.color.rgb = accent_gold
        shape.line.width = Pt(2)
        
        tf = shape.text_frame
        tf.word_wrap = True
        p1 = tf.paragraphs[0]
        p1.text = f"{title}\n"
        p1.font.bold = True
        p1.font.size = Pt(18)
        p1.font.color.rgb = accent_gold
        p1.alignment = PP_ALIGN.LEFT
        
        p2 = tf.add_paragraph()
        p2.text = desc
        p2.font.size = Pt(14)
        p2.font.color.rgb = white
        p2.alignment = PP_ALIGN.LEFT
        
    slide.shapes.add_connector(1, Inches(4.2), Inches(3.9), Inches(5), Inches(2.9))
    slide.shapes.add_connector(1, Inches(8.2), Inches(2.9), Inches(9), Inches(3.9))
    slide.shapes.add_connector(1, Inches(9.1), Inches(4.8), Inches(8.2), Inches(5.9))

    # ==========================================
    # Slide 6: Core Process 2 - Scouting
    # ==========================================
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_bg(slide, bg_light)
    add_title(slide, "Core Process 2: Scouting & Sponsoring", bg_dark)

    steps2 = [
        ("Search", "Sponsor filters grid by category.", Inches(1)),
        ("Review", "Analyzes the Verified Student Profile.", Inches(5)),
        ("Dispatch", "Sends formal scholarship offer via platform.", Inches(9))
    ]
    
    for i, (title, desc, left) in enumerate(steps2):
        shape = slide.shapes.add_shape(MSO_SHAPE.HEXAGON, left, Inches(3.5), Inches(3.2), Inches(2.5))
        shape.fill.solid()
        shape.fill.fore_color.rgb = white
        shape.line.color.rgb = bg_dark
        shape.line.width = Pt(3)
        
        tf = shape.text_frame
        tf.word_wrap = True
        p1 = tf.paragraphs[0]
        p1.text = f"\n{title}\n"
        p1.font.bold = True
        p1.font.size = Pt(22)
        p1.font.color.rgb = bg_dark
        p1.alignment = PP_ALIGN.CENTER
        
        p2 = tf.add_paragraph()
        p2.text = desc
        p2.font.size = Pt(16)
        p2.font.color.rgb = text_dark
        p2.alignment = PP_ALIGN.CENTER
        
        if i < 2:
            arrow = slide.shapes.add_shape(MSO_SHAPE.CHEVRON, left + Inches(3.3), Inches(4.3), Inches(0.8), Inches(0.8))
            arrow.fill.solid()
            arrow.fill.fore_color.rgb = accent_gold

    # ==========================================
    # Slide 7: Prototype UX
    # ==========================================
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_bg(slide, bg_dark)
    add_title(slide, "Prototype Description (UI/UX)", white)
    
    features = [
        ("Institutional Glassmorphism", "Translucent navbars, soft rounded corners, and dynamic multi-layered drop shadows representing professional, trustworthy corporate software."),
        ("Role-Based Workspaces", "Scouts, Referees, and Sponsors dynamically switch views within a fast Single Page Application (SPA) structure."),
        ("Data-Driven Grid", "Profiles rendered instantly via modular JavaScript with powerful filtering mechanisms.")
    ]
    
    for i, (title, desc) in enumerate(features):
        shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1), Inches(2 + (i*1.8)), Inches(11.33), Inches(1.5))
        shape.fill.solid()
        shape.fill.fore_color.rgb = bg_light
        
        tf = shape.text_frame
        tf.word_wrap = True
        p1 = tf.paragraphs[0]
        p1.text = f"{title}"
        p1.font.bold = True
        p1.font.size = Pt(20)
        p1.font.color.rgb = bg_dark
        
        p2 = tf.add_paragraph()
        p2.text = f"\n{desc}"
        p2.font.size = Pt(16)
        p2.font.color.rgb = text_dark

    # ==========================================
    # Slide 8: Future Roadmap
    # ==========================================
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_bg(slide, bg_light)
    add_title(slide, "Beyond the MVP: Roadmap", bg_dark)

    phases = [
        ("Phase 2", "Video Analytics", "Deep performance tracking powered by AI video review."),
        ("Phase 3", "Blockchain Ledger", "100% tamper-proof immutable records for universities."),
        ("Phase 4", "Mobile Student App", "Native app for students to carry their profiles globally.")
    ]
    
    for i, (phase, title, desc) in enumerate(phases):
        top = Inches(2.5 + (i * 1.5))
        left = Inches(1 + (i * 2))
        
        circle = slide.shapes.add_shape(MSO_SHAPE.OVAL, left, top, Inches(0.8), Inches(0.8))
        circle.fill.solid()
        circle.fill.fore_color.rgb = accent_gold
        circle.line.fill.background()
        circle.text_frame.paragraphs[0].text = f"{i+2}"
        circle.text_frame.paragraphs[0].font.bold = True
        circle.text_frame.paragraphs[0].alignment = PP_ALIGN.CENTER
        
        shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left + Inches(1), top, Inches(6), Inches(1.2))
        shape.fill.solid()
        shape.fill.fore_color.rgb = bg_dark
        
        tf = shape.text_frame
        tf.word_wrap = True
        p1 = tf.paragraphs[0]
        p1.text = f"{phase}: {title}"
        p1.font.bold = True
        p1.font.size = Pt(18)
        p1.font.color.rgb = accent_gold
        
        p2 = tf.add_paragraph()
        p2.text = f"{desc}"
        p2.font.size = Pt(14)
        p2.font.color.rgb = white

    # ==========================================
    # Slide 9: Market Opportunity (TAM)
    # ==========================================
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_bg(slide, bg_light)
    add_title(slide, "Market Opportunity (TAM)", bg_dark)

    shape = slide.shapes.add_shape(MSO_SHAPE.DONUT, Inches(2), Inches(2.5), Inches(3.5), Inches(3.5))
    shape.fill.solid()
    shape.fill.fore_color.rgb = bg_dark
    shape.line.color.rgb = accent_gold
    
    tf = shape.text_frame
    p1 = tf.paragraphs[0]
    p1.text = "$12B"
    p1.font.size = Pt(36)
    p1.font.bold = True
    p1.font.color.rgb = white
    p1.alignment = PP_ALIGN.CENTER

    shape_text = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.5), Inches(2.5), Inches(5.5), Inches(3.5))
    shape_text.fill.solid()
    shape_text.fill.fore_color.rgb = white
    shape_text.line.color.rgb = bg_dark
    
    tf2 = shape_text.text_frame
    tf2.word_wrap = True
    
    p2 = tf2.paragraphs[0]
    p2.text = "The Global Opportunity\n\n"
    p2.font.bold = True
    p2.font.size = Pt(24)
    p2.font.color.rgb = bg_dark
    
    p3 = tf2.add_paragraph()
    p3.text = "• Youth Sports & Arts Scholarships represent a multi-billion dollar industry globally.\n"
    p3.font.size = Pt(18)
    p3.font.color.rgb = text_dark
    
    p4 = tf2.add_paragraph()
    p4.text = "• Thousands of foundations lack a streamlined tool to discover verified, eligible candidates at scale."
    p4.font.size = Pt(18)
    p4.font.color.rgb = text_dark

    # ==========================================
    # Slide 10: Competitive Advantage (USP)
    # ==========================================
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_bg(slide, bg_dark)
    add_title(slide, "Our Unique Selling Proposition (USP)", white)

    usps = [
        ("Current Systems", "TalentLink", Inches(2), bg_light, bg_dark, text_dark),
        ("Fragmented Spreadsheets", "Unified Database", Inches(3.5), bg_dark, accent_gold, white),
        ("Self-Reported / Unverified", "100% Certified Data", Inches(4.5), bg_dark, accent_gold, white),
        ("Manual Email Inquiries", "1-Click Direct Offers", Inches(5.5), bg_dark, accent_gold, white)
    ]
    
    for left_text, right_text, top, bg_color, left_txt_col, right_txt_col in usps:
        # Left Box (Status Quo)
        shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(1.5), top, Inches(4), Inches(0.8))
        shape.fill.solid()
        shape.fill.fore_color.rgb = bg_light if top == Inches(2) else white
        shape.line.fill.background()
        
        p1 = shape.text_frame.paragraphs[0]
        p1.text = left_text
        p1.font.bold = top == Inches(2)
        p1.font.size = Pt(18) if top == Inches(2) else Pt(16)
        p1.font.color.rgb = bg_dark if top == Inches(2) else text_dark
        p1.alignment = PP_ALIGN.CENTER
        
        # Middle Arrow
        if top != Inches(2):
            arrow = slide.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, Inches(6), top + Inches(0.2), Inches(0.6), Inches(0.4))
            arrow.fill.solid()
            arrow.fill.fore_color.rgb = accent_gold
            arrow.line.fill.background()

        # Right Box (TalentLink)
        shape2 = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(7), top, Inches(5), Inches(0.8))
        shape2.fill.solid()
        shape2.fill.fore_color.rgb = bg_color
        shape2.line.color.rgb = accent_gold if top != Inches(2) else white
        
        p2 = shape2.text_frame.paragraphs[0]
        p2.text = right_text
        p2.font.bold = top == Inches(2)
        p2.font.size = Pt(18) if top == Inches(2) else Pt(18)
        p2.font.color.rgb = left_txt_col
        p2.alignment = PP_ALIGN.CENTER

    # ==========================================
    # Slide 11: Call to Action (The Ask)
    # ==========================================
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_bg(slide, bg_light)
    add_title(slide, "The Ask (Call to Action)", bg_dark)

    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(2), Inches(2.5), Inches(9.33), Inches(3.5))
    shape.fill.solid()
    shape.fill.fore_color.rgb = bg_dark
    shape.line.color.rgb = accent_gold
    shape.line.width = Pt(3)
    
    tf = shape.text_frame
    tf.word_wrap = True
    
    p1 = tf.paragraphs[0]
    p1.text = "\nWe are seeking Pilot Partners\n"
    p1.font.bold = True
    p1.font.size = Pt(36)
    p1.font.color.rgb = accent_gold
    p1.alignment = PP_ALIGN.CENTER
    
    p2 = tf.add_paragraph()
    p2.text = "Join us in launching the MVP across 5 core educational institutions."
    p2.font.size = Pt(22)
    p2.font.color.rgb = white
    p2.alignment = PP_ALIGN.CENTER

    # ==========================================
    # Slide 12: Q&A / Contact
    # ==========================================
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_bg(slide, bg_dark)

    txBox = slide.shapes.add_textbox(Inches(1), Inches(2.5), Inches(11.33), Inches(1.5))
    p = txBox.text_frame.paragraphs[0]
    p.text = "Thank You. Questions?"
    p.font.size = Pt(64)
    p.font.bold = True
    p.font.color.rgb = white
    p.alignment = PP_ALIGN.CENTER
    
    txBox2 = slide.shapes.add_textbox(Inches(1), Inches(4.5), Inches(11.33), Inches(1))
    p2 = txBox2.text_frame.paragraphs[0]
    p2.text = "contact@talentlink.edu  |  www.talentlink.edu"
    p2.font.size = Pt(24)
    p2.font.color.rgb = accent_gold
    p2.alignment = PP_ALIGN.CENTER

    prs.save('TalentLink_Presentation_v4.pptx')
    print("Presentation created successfully as TalentLink_Presentation_v4.pptx")

if __name__ == '__main__':
    create_rich_presentation()
